from io import BytesIO
from unittest import mock

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from docx import Document
from rest_framework import status
from rest_framework.test import APITestCase

from accounts.models import User
from courses.content_analyzer import ContentAnalysisError, analyze_upload
from courses.demo_files import build_pdf
from courses.models import Category, CourseTranscript

from .utils import (
    MP4_BYTES,
    TemporaryMediaMixin,
    create_course,
    pdf_upload,
    wav_upload,
)

ANALYZE_URL = "/api/courses/analyze-content/"

PROGRAMMING_TEXT = (
    "Welcome to Python Programming Basics. This beginner course explains "
    "variables, loops, functions and simple algorithms. You will write code, "
    "debug scripts and build a small project."
)


def rich_pdf_upload():
    content = build_pdf(
        "Python Programming Basics",
        "Learn variables, loops and functions in Python. This beginner course "
        "covers coding fundamentals and simple algorithms.",
    )
    return SimpleUploadedFile(
        "python_basics.pdf", content, content_type="application/pdf"
    )


def docx_upload():
    document = Document()
    document.core_properties.title = "Design Fundamentals"
    document.add_heading("Design Fundamentals", level=1)
    document.add_paragraph(
        "Learn typography, colour, layout and accessibility. This beginner "
        "course teaches visual design with Figma and practical prototypes."
    )
    buffer = BytesIO()
    document.save(buffer)
    return SimpleUploadedFile(
        "design_fundamentals.docx",
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
    )


@override_settings(
    PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"],
)
class ContentAnalysisApiTests(TemporaryMediaMixin, APITestCase):
    password = "StrongPass123!"

    def setUp(self):
        self.teacher = User.objects.create_user(
            email="teacher@example.com",
            password=self.password,
            first_name="Thomas",
            last_name="Teacher",
            role=User.Role.TEACHER,
        )
        self.other_teacher = User.objects.create_user(
            email="other@example.com",
            password=self.password,
            first_name="Olivia",
            last_name="Other",
            role=User.Role.TEACHER,
        )
        self.student = User.objects.create_user(
            email="student@example.com",
            password=self.password,
            first_name="Sam",
            last_name="Student",
            role=User.Role.STUDENT,
        )
        self.programming = Category.objects.create(
            name="Programming",
            description="Python, JavaScript, algorithms and software development.",
        )
        self.design = Category.objects.create(
            name="Design",
            description="Typography, colour, layout, Figma and visual design.",
        )
        self.client.force_login(self.teacher)

    def test_teacher_can_analyze_an_uploaded_pdf(self):
        response = self.client.post(
            ANALYZE_URL,
            {"file": rich_pdf_upload()},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.json()
        self.assertEqual(data["title"], "Python Programming Basics")
        self.assertEqual(data["category"], "Programming")
        self.assertEqual(data["category_id"], self.programming.pk)
        self.assertEqual(data["difficulty"], "beginner")
        self.assertEqual(data["source"], "pdf")
        self.assertGreaterEqual(len(data["tags"]), 5)
        self.assertLessEqual(len(data["tags"]), 8)
        self.assertIn("python", data["tags"])
        self.assertGreaterEqual(len(data["description"]), 150)
        self.assertLessEqual(len(data["description"]), 300)
        self.assertGreaterEqual(data["processing_time_ms"], 0)
        for value in data["confidence"].values():
            self.assertGreaterEqual(value, 0)
            self.assertLessEqual(value, 1)

    def test_teacher_can_analyze_a_docx_document(self):
        response = self.client.post(
            ANALYZE_URL,
            {"file": docx_upload()},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.json()
        self.assertEqual(data["title"], "Design Fundamentals")
        self.assertEqual(data["category"], "Design")
        self.assertEqual(data["source"], "docx")
        self.assertIn("typography", data["tags"])

    @override_settings(CONTENT_ANALYZER_WHISPER_MODEL="tiny")
    def test_audio_is_transcribed_locally(self):
        speech = (
            "Welcome to this machine learning course. In this lesson we "
            "introduce neural networks, gradient descent and model training. "
            "You will build a simple classifier with Python and scikit-learn."
        )
        with mock.patch(
            "courses.content_analyzer.transcribe_audio",
            return_value=speech,
        ):
            response = self.client.post(
                ANALYZE_URL,
                {"file": wav_upload()},
                format="multipart",
            )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.json()
        self.assertEqual(data["source"], "audio")
        self.assertEqual(data["title"], "Audio")
        self.assertIn("neural", data["tags"])
        self.assertEqual(data["difficulty"], "advanced")

    def test_media_transcription_failure_degrades_gracefully(self):
        with mock.patch(
            "courses.content_analyzer.transcribe_audio",
            side_effect=ContentAnalysisError(
                "Local speech recognition is not installed on this server.",
                code="transcription_unavailable",
            ),
        ):
            response = self.client.post(
                ANALYZE_URL,
                {"file": wav_upload()},
                format="multipart",
            )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.json()
        self.assertEqual(data["source"], "audio")
        self.assertEqual(data["title"], "Audio")
        self.assertTrue(data["warnings"])
        self.assertEqual(data["word_count"], 0)

    def test_existing_course_can_be_analyzed_from_its_transcript(self):
        course = create_course(
            self.teacher,
            title="Neural networks",
            filename="lesson.mp4",
            content=MP4_BYTES,
        )
        CourseTranscript.objects.create(
            course=course,
            status=CourseTranscript.Status.COMPLETED,
            text=PROGRAMMING_TEXT,
        )

        response = self.client.post(
            ANALYZE_URL,
            {"course_id": course.pk},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.json()
        self.assertEqual(data["source"], "transcript")
        self.assertEqual(data["title"], "Neural networks")
        self.assertEqual(data["category"], "Programming")

    def test_existing_course_file_can_be_analyzed(self):
        course = create_course(
            self.teacher,
            title="Python course",
            description="Coding fundamentals",
            filename="course.pdf",
            content=build_pdf("Python Programming Basics", PROGRAMMING_TEXT),
        )

        response = self.client.post(
            ANALYZE_URL,
            {"course_id": course.pk},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.json()
        self.assertEqual(data["source"], "pdf")
        self.assertEqual(data["category"], "Programming")

    def test_students_cannot_analyze_content(self):
        self.client.force_login(self.student)

        response = self.client.post(
            ANALYZE_URL,
            {"file": rich_pdf_upload()},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_teachers_cannot_analyze_another_teachers_course(self):
        course = create_course(self.other_teacher, title="Other course")

        response = self.client.post(
            ANALYZE_URL,
            {"course_id": course.pk},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_exactly_one_input_source_is_required(self):
        response = self.client.post(ANALYZE_URL, {}, format="multipart")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

        response = self.client.post(
            ANALYZE_URL,
            {"file": rich_pdf_upload(), "course_id": 1},
            format="multipart",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_unsupported_files_are_rejected(self):
        upload = SimpleUploadedFile(
            "notes.txt", b"plain text", content_type="text/plain"
        )

        response = self.client.post(
            ANALYZE_URL,
            {"file": upload},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("file", response.json()["error"]["fields"])

    def test_anonymous_users_are_rejected(self):
        self.client.logout()

        response = self.client.post(
            ANALYZE_URL,
            {"file": rich_pdf_upload()},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_local_analyzer_rejects_empty_uploads(self):
        upload = SimpleUploadedFile(
            "empty.pdf", b"", content_type="application/pdf"
        )

        with self.assertRaises(ContentAnalysisError):
            analyze_upload(upload)

    def test_difficulty_is_persisted_and_filterable(self):
        course = create_course(self.teacher, title="Advanced topic")
        course.difficulty = "advanced"
        course.save(update_fields=["difficulty"])

        response = self.client.get("/api/courses/?difficulty=advanced")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        results = response.json()["results"]
        self.assertEqual([item["id"] for item in results], [course.pk])
        self.assertEqual(results[0]["difficulty"], "advanced")
        self.assertEqual(results[0]["difficulty_display"], "Advanced")

    def test_difficulty_is_validated(self):
        response = self.client.get("/api/courses/?difficulty=expert")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("difficulty", response.json()["error"]["fields"])

    def test_course_creation_accepts_a_difficulty(self):
        response = self.client.post(
            "/api/courses/",
            {
                "title": "Beginner course",
                "file": pdf_upload(),
                "difficulty": "beginner",
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.json()["difficulty"], "beginner")
        self.assertEqual(response.json()["difficulty_display"], "Beginner")
