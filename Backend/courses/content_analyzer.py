"""Local content analysis used by the course form auto-fill.

Documents are read with pypdf and python-docx, media is decoded with the
FFmpeg binary bundled by imageio-ffmpeg, speech is transcribed with a local
OpenAI Whisper model, and the remaining NLP is scikit-learn and deterministic
rules. No external API is called.
"""

from __future__ import annotations

import logging
import re
import shutil
import subprocess
import tempfile
import threading
import time
import wave
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import imageio_ffmpeg
import numpy as np
from django.conf import settings
from django.core.exceptions import ValidationError as DjangoValidationError
from sklearn.feature_extraction.text import ENGLISH_STOP_WORDS, TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from .models import Category, CourseTranscript
from .validators import (
    AUDIO_EXTENSIONS,
    PDF_EXTENSIONS,
    VIDEO_EXTENSIONS,
    validate_course_file,
)

logger = logging.getLogger(__name__)

ANALYZABLE_EXTENSIONS = PDF_EXTENSIONS + AUDIO_EXTENSIONS + VIDEO_EXTENSIONS + (
    "docx",
)
DEFAULT_MAX_WORDS = 5000
DEFAULT_PDF_PAGES = 5
DEFAULT_MEDIA_SECONDS = 180
DEFAULT_WHISPER_MODEL = "tiny"
DOCX_MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024

TITLE_MIN_WORDS = 2
TITLE_MAX_WORDS = 14
DESCRIPTION_MAX_CHARS = 300
DESCRIPTION_MIN_CHARS = 150
READING_WORDS_PER_MINUTE = 200
DIFFICULTY_ORDER = {"beginner": 0, "intermediate": 1, "advanced": 2}

MEDIA_DURATION_PATTERN = re.compile(
    r"Duration:\s*(\d+):(\d{2}):(\d{2}(?:\.\d+)?)"
)
FILENAME_NUMBER_PREFIX = re.compile(r"^\s*(\d+)")

TOKEN_PATTERN = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ][A-Za-zÀ-ÖØ-öø-ÿ'’-]*")
SENTENCE_PATTERN = re.compile(r"(?<=[.!?])\s+(?=[A-Z0-9À-ÖØ-Þ])")
HEX_NAME_PATTERN = re.compile(r"\b[0-9a-f]{16,}\b", re.IGNORECASE)

GENERIC_TAGS = {
    "course",
    "lesson",
    "module",
    "chapter",
    "video",
    "audio",
    "document",
    "content",
    "cover",
    "covers",
    "student",
    "students",
    "teacher",
    "learn",
    "learning",
    "understand",
    "introduction",
    "overview",
    "summary",
    "welcome",
    "hello",
    "today",
    "example",
    "examples",
    "using",
    "used",
    "will",
    "also",
    "can",
    "may",
}

STOPWORDS = set(ENGLISH_STOP_WORDS) | GENERIC_TAGS

ACRONYMS = {
    "ai": "AI",
    "api": "API",
    "css": "CSS",
    "html": "HTML",
    "http": "HTTP",
    "javascript": "JavaScript",
    "ml": "ML",
    "pdf": "PDF",
    "sql": "SQL",
    "ui": "UI",
    "ux": "UX",
}

CATEGORY_KEYWORDS = {
    "programming": {
        "algorithm",
        "api",
        "backend",
        "class",
        "code",
        "coding",
        "compiler",
        "developer",
        "development",
        "debug",
        "debugging",
        "django",
        "frontend",
        "function",
        "git",
        "java",
        "javascript",
        "loop",
        "object",
        "programming",
        "python",
        "react",
        "recursion",
        "repository",
        "script",
        "software",
        "syntax",
        "typescript",
        "variable",
    },
    "web": {
        "angular",
        "api",
        "backend",
        "browser",
        "css",
        "deployment",
        "django",
        "flask",
        "frontend",
        "html",
        "http",
        "javascript",
        "node",
        "react",
        "responsive",
        "server",
        "ui",
        "ux",
        "vue",
        "web",
        "website",
    },
    "data": {
        "analytics",
        "classification",
        "clustering",
        "data",
        "database",
        "dataset",
        "machine learning",
        "metrics",
        "model",
        "numpy",
        "pandas",
        "probability",
        "python",
        "regression",
        "sql",
        "statistics",
        "tableau",
        "visualization",
    },
    "machine learning": {
        "classification",
        "clustering",
        "dataset",
        "deep learning",
        "gradient",
        "machine learning",
        "model",
        "neural",
        "numpy",
        "pandas",
        "python",
        "regression",
        "tensor",
        "training",
    },
    "design": {
        "accessibility",
        "branding",
        "colour",
        "color",
        "design",
        "figma",
        "graphic",
        "layout",
        "prototype",
        "sketch",
        "typography",
        "ui",
        "usability",
        "ux",
        "visual",
        "wireframe",
    },
    "business": {
        "accounting",
        "business",
        "customer",
        "economics",
        "entrepreneurship",
        "finance",
        "leadership",
        "management",
        "marketing",
        "operations",
        "product",
        "project",
        "sales",
        "startup",
        "strategy",
    },
    "language": {
        "conversation",
        "english",
        "fluency",
        "french",
        "german",
        "grammar",
        "language",
        "listening",
        "pronunciation",
        "reading",
        "spanish",
        "speaking",
        "translation",
        "vocabulary",
        "writing",
    },
    "math": {
        "algebra",
        "calculus",
        "derivative",
        "equation",
        "geometry",
        "integral",
        "math",
        "mathematics",
        "matrix",
        "number",
        "probability",
        "proof",
        "statistics",
        "theorem",
        "vector",
    },
    "science": {
        "astronomy",
        "atom",
        "biology",
        "cell",
        "chemistry",
        "energy",
        "experiment",
        "genetics",
        "hypothesis",
        "laboratory",
        "molecule",
        "physics",
        "research",
        "science",
        "theory",
    },
    "music": {
        "audio",
        "chord",
        "composition",
        "guitar",
        "harmony",
        "instrument",
        "melody",
        "music",
        "piano",
        "rhythm",
        "scale",
        "singing",
        "song",
        "tempo",
    },
    "photography": {
        "aperture",
        "camera",
        "composition",
        "editing",
        "exposure",
        "image",
        "lens",
        "lighting",
        "photography",
        "photoshop",
        "portrait",
        "shutter",
    },
    "health": {
        "anatomy",
        "clinical",
        "exercise",
        "fitness",
        "health",
        "medical",
        "medicine",
        "mental",
        "nutrition",
        "nursing",
        "patient",
        "physiology",
        "wellness",
    },
}

TECHNICAL_TERMS = {
    "algorithm",
    "api",
    "architecture",
    "array",
    "asynchronous",
    "backend",
    "binary",
    "class",
    "client",
    "compiler",
    "complexity",
    "concurrency",
    "database",
    "dataset",
    "derivative",
    "distributed",
    "equation",
    "framework",
    "function",
    "gradient",
    "http",
    "integral",
    "interface",
    "kubernetes",
    "matrix",
    "method",
    "model",
    "neural",
    "object",
    "optimization",
    "parameter",
    "polymorphism",
    "probability",
    "protocol",
    "react",
    "recursion",
    "regression",
    "server",
    "sql",
    "statistics",
    "tensor",
    "theorem",
    "thread",
    "vector",
    "webhook",
}

ADVANCED_TERMS = {
    "advanced",
    "architecture",
    "asynchronous",
    "concurrency",
    "distributed",
    "eigenvalue",
    "gradient",
    "inference",
    "metaprogramming",
    "optimization",
    "polymorphism",
    "quantum",
    "recursion",
    "scalability",
    "stochastic",
    "tensor",
    "topology",
}

BEGINNER_TERMS = {
    "basics",
    "beginner",
    "essential",
    "essentials",
    "first",
    "fundamental",
    "fundamentals",
    "getting started",
    "intro",
    "introduction",
    "overview",
    "simple",
    "starter",
}


class ContentAnalysisError(Exception):
    """Raised when content cannot be analyzed for a reason safe to show."""

    def __init__(self, message, code="analysis_failed"):
        super().__init__(message)
        self.code = code


@dataclass
class ExtractedContent:
    text: str = ""
    title_hint: str | None = None
    source: str = "text"
    warnings: list[str] = field(default_factory=list)
    word_count: int = 0


_WHISPER_MODEL = None
_WHISPER_LOCK = threading.Lock()


def _setting(name, default):
    return getattr(settings, name, default)


def _max_words():
    value = _setting("CONTENT_ANALYZER_MAX_WORDS", DEFAULT_MAX_WORDS)
    return max(200, int(value))


def _pdf_pages():
    value = _setting("CONTENT_ANALYZER_PDF_PAGES", DEFAULT_PDF_PAGES)
    return max(1, int(value))


def _media_seconds():
    value = _setting("CONTENT_ANALYZER_MEDIA_SECONDS", DEFAULT_MEDIA_SECONDS)
    return max(10, int(value))


def _clean_text(text):
    text = (text or "").replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _cap_words(text, limit=None):
    limit = limit or _max_words()
    words = text.split()
    if len(words) <= limit:
        return text
    return " ".join(words[:limit])


def _split_sentences(text):
    normalized = re.sub(r"\s+", " ", text or "").strip()
    if not normalized:
        return []

    sentences = [
        sentence.strip()
        for sentence in SENTENCE_PATTERN.split(normalized)
        if sentence.strip()
    ]
    if len(sentences) <= 1:
        words = normalized.split()
        if len(words) > 35:
            sentences = [
                " ".join(words[index : index + 22])
                for index in range(0, min(len(words), 110), 22)
            ]

    cleaned = []
    for sentence in sentences:
        sentence = sentence.strip(" -•\t")
        if len(sentence) < 20:
            continue
        if len(sentence) > 400:
            sentence = sentence[:400].rsplit(" ", 1)[0]
        cleaned.append(sentence)
    return cleaned


def _terminate(sentence):
    sentence = sentence.rstrip()
    if not sentence:
        return sentence
    if sentence[-1] not in ".!?":
        sentence += "."
    return sentence


def _smart_title(value):
    words = []
    for word in value.split():
        lowered = word.lower()
        words.append(ACRONYMS.get(lowered, word[:1].upper() + word[1:]))
    return " ".join(words)


def _title_from_filename(filename):
    stem = Path(filename or "").stem
    stem = HEX_NAME_PATTERN.sub(" ", stem)
    stem = re.sub(r"[_\-.]+", " ", stem)
    stem = re.sub(r"\b(copy|final|draft|version|v\d+|new)\b", " ", stem, flags=re.I)
    stem = re.sub(r"^\s*\d+\s*", "", stem)
    stem = re.sub(r"\s+", " ", stem).strip()
    if not stem:
        return "Untitled course", 0.3
    words = stem.split()
    if len(words) > TITLE_MAX_WORDS:
        stem = " ".join(words[:TITLE_MAX_WORDS])
    title = _smart_title(stem)
    confidence = 0.55 if len(title) >= 12 else 0.4
    return title[:200], confidence


def _valid_heading(line):
    words = TOKEN_PATTERN.findall(line)
    if not TITLE_MIN_WORDS <= len(words) <= TITLE_MAX_WORDS:
        return False
    if not 12 <= len(line) <= 120:
        return False
    if "://" in line or "@" in line:
        return False
    if sum(character.isdigit() for character in line) > 4:
        return False
    if line.endswith((".", ",", ";")):
        return False
    if len(words) > 8 and line[-1] in "!?":
        return False
    return True


def _title_from_heading(text):
    lines = [line.strip() for line in (text or "").splitlines() if line.strip()]
    best = None
    best_score = -1.0

    for index, line in enumerate(lines[:15]):
        if not _valid_heading(line):
            continue
        words = line.split()
        score = 0.0
        if 3 <= len(words) <= 10:
            score += 2.0
        score += max(0.0, 2.0 - index * 0.25)
        score += 0.5 if line.istitle() else 0.0
        score += min(1.0, len(line) / 80)
        if score > best_score:
            best = line
            best_score = score

    if best is None:
        return None, 0.0
    return best[:200], 0.82 if best_score >= 3 else 0.68


def _guess_title(text, title_hint, filename):
    hint = _clean_text(title_hint or "")
    if hint and 3 <= len(hint) <= 200:
        return hint[:200], 0.95

    heading, heading_confidence = _title_from_heading(text)
    if heading:
        return heading, heading_confidence

    return _title_from_filename(filename)


def _strip_heading_lines(text, title):
    title_key = re.sub(r"\s+", " ", title or "").lower().strip(" .:-—")
    output = []
    header_done = False

    for line in (text or "").splitlines():
        stripped = line.strip()
        if not stripped:
            output.append("")
            continue

        if not header_done:
            normalized = re.sub(r"\s+", " ", stripped).lower().strip(" .:-—")
            if normalized == title_key:
                header_done = True
                continue
            if (
                len(TOKEN_PATTERN.findall(stripped)) <= 2
                and len(stripped) <= 30
                and not output
            ):
                continue
            header_done = True

        output.append(line)

    result = "\n".join(output).strip()
    return result or text


def _category_hints(name):
    normalized = f" {name.lower()} "
    hints = set()
    for keyword, keyword_hints in CATEGORY_KEYWORDS.items():
        if f" {keyword} " in normalized or keyword in normalized:
            hints.update(keyword_hints)

    expanded = set()
    for hint in hints:
        expanded.add(hint)
        if " " in hint:
            expanded.update(hint.split())

    expanded.update(
        token
        for token in TOKEN_PATTERN.findall(normalized.lower())
        if token not in STOPWORDS and len(token) >= 3
    )
    return expanded


def _token_counts(text):
    tokens = [
        token.lower().strip("'’-")
        for token in TOKEN_PATTERN.findall(text or "")
    ]
    return [token for token in tokens if token]


def _summary_fallback(title, tags):
    if tags:
        return f"This course explores {title.lower()} and covers {', '.join(tags[:4])}."
    return f"This course introduces the key ideas behind {title.lower()}."


def _summarize(text, title, tags):
    sentences = _split_sentences(text)
    if not sentences:
        return _terminate(_summary_fallback(title, tags))

    if len(sentences) == 1:
        summary = sentences[0]
    else:
        try:
            vectorizer = TfidfVectorizer(
                stop_words="english",
                ngram_range=(1, 2),
                max_features=600,
                sublinear_tf=True,
            )
            matrix = vectorizer.fit_transform(sentences)
            similarity = cosine_similarity(matrix)
            np.fill_diagonal(similarity, 0.0)
            row_sums = similarity.sum(axis=1)
            row_sums[row_sums == 0] = 1.0
            transition = similarity / row_sums[:, None]

            scores = np.full(len(sentences), 1.0 / len(sentences))
            for _ in range(30):
                scores = 0.15 / len(sentences) + 0.85 * transition.T.dot(scores)
        except ValueError:
            scores = np.ones(len(sentences))

        for index in range(len(sentences)):
            scores[index] *= 1.0 / (1.0 + index * 0.25)
            word_count = len(sentences[index].split())
            if word_count < 6:
                scores[index] *= 0.6
            elif word_count > 45:
                scores[index] *= 0.75

        chosen = sorted(
            np.argsort(scores)[-3:].tolist(),
            key=lambda index: index,
        )
        summary = " ".join(sentences[index] for index in chosen)

    summary = _terminate(re.sub(r"\s+", " ", summary).strip())

    if len(summary) > DESCRIPTION_MAX_CHARS:
        summary = summary[:DESCRIPTION_MAX_CHARS].rsplit(" ", 1)[0]
        summary = _terminate(summary)

    if len(summary) < DESCRIPTION_MIN_CHARS:
        extra = _terminate(_summary_fallback(title, tags))
        if len(summary) + 1 + len(extra) <= DESCRIPTION_MAX_CHARS:
            summary = f"{summary} {extra}"
        elif len(summary) < DESCRIPTION_MIN_CHARS and tags:
            extra = f"It covers {', '.join(tags[:4])}."
            if len(summary) + 1 + len(extra) <= DESCRIPTION_MAX_CHARS:
                summary = f"{summary} {extra}"

    return summary


def _predict_category(text, title, tags, categories):
    if not categories:
        return None, None, 0.0

    document = _cap_words(f"{title}. {text}. {' '.join(tags)}", 1500)
    if not document.strip():
        return None, None, 0.0

    documents = [document]
    for category in categories:
        hints = " ".join(sorted(_category_hints(category.name)))
        documents.append(
            f"{category.name} {category.name} {category.description} {hints}"
        )

    try:
        vectorizer = TfidfVectorizer(
            stop_words="english",
            ngram_range=(1, 2),
            sublinear_tf=True,
        )
        matrix = vectorizer.fit_transform(documents)
        similarities = cosine_similarity(matrix[0:1], matrix[1:]).ravel()
    except ValueError:
        similarities = np.zeros(len(categories))

    tokens = set(_token_counts(document))
    scores = []
    for index, category in enumerate(categories):
        overlap = len(tokens & _category_hints(category.name))
        scores.append(float(similarities[index]) + overlap * 0.06)

    order = sorted(range(len(categories)), key=lambda index: scores[index], reverse=True)
    best_index = order[0]
    best_score = scores[best_index]
    if best_score <= 0:
        return None, None, 0.0

    runner_up = scores[order[1]] if len(order) > 1 else 0.0
    margin = max(0.0, best_score - runner_up)
    confidence = 0.45 + min(0.42, margin * 2.0) + min(0.08, best_score * 0.15)
    confidence = round(min(0.97, confidence), 2)
    category = categories[best_index]
    return category.name, category.pk, confidence


def _valid_tag(term):
    if not term or len(term) < 3 or len(term) > 50:
        return False
    if term.isdigit():
        return False
    words = term.split()
    if not words or words[0] in STOPWORDS or words[-1] in STOPWORDS:
        return False
    if any(len(word) < 2 for word in words):
        return False
    return True


def _extract_tags(text, title, category_name=None, limit=8):
    combined = _cap_words(f"{title}. {text}", 3000)
    tokens = _token_counts(combined)
    if not tokens:
        return []

    counts = Counter(tokens)
    sentences = _split_sentences(combined)
    tfidf_scores = {}

    if sentences:
        try:
            vectorizer = TfidfVectorizer(
                stop_words="english",
                ngram_range=(1, 2),
                min_df=1,
                sublinear_tf=True,
            )
            matrix = vectorizer.fit_transform(sentences)
            if matrix.shape[0] > 1:
                mean_scores = np.asarray(matrix.mean(axis=0)).ravel()
                tfidf_scores = dict(
                    zip(vectorizer.get_feature_names_out(), mean_scores)
                )
            else:
                tfidf_scores = {
                    term: 1.0 for term in vectorizer.get_feature_names_out()
                }
        except ValueError:
            tfidf_scores = {}

    domain_hints = set()
    if category_name:
        domain_hints = _category_hints(category_name)
    for hints in CATEGORY_KEYWORDS.values():
        domain_hints.update(hints)

    candidates = set(tfidf_scores)
    candidates.update(counts)
    for first, second in zip(tokens, tokens[1:]):
        if first in STOPWORDS or second in STOPWORDS:
            continue
        candidates.add(f"{first} {second}")

    ranked = []
    for term in candidates:
        if not _valid_tag(term):
            continue
        count = counts.get(term, 0)
        if " " in term:
            count = max(1, counts.get(term.split()[0], 0))
        score = tfidf_scores.get(term, 0.0)
        score += min(0.5, np.log1p(count) * 0.12)
        if " " in term:
            score += 0.12
        if term in domain_hints:
            score += 0.35
        if term in title.lower():
            score += 0.05
        ranked.append((term, float(score)))

    ranked.sort(key=lambda item: (-item[1], item[0]))

    tags = []
    seen_words = set()
    for term, _score in ranked:
        words = set(term.split())
        if " " in term and words & seen_words:
            continue
        if term in seen_words:
            continue
        tags.append(term)
        seen_words.update(words if " " in term else {term})
        if len(tags) >= limit:
            break

    if category_name and len(tags) < limit:
        lowered = category_name.lower()
        if lowered not in tags and _valid_tag(lowered):
            tags.append(lowered)

    if len(tags) < 3:
        for word in _token_counts(title):
            if word in tags or word in STOPWORDS or len(word) < 3:
                continue
            tags.append(word)
            if len(tags) >= limit:
                break

    return tags[:limit]


def _classify_difficulty(text, title):
    combined = _cap_words(f"{title}. {text}", 3000)
    tokens = _token_counts(combined)
    if len(tokens) < 12:
        return "beginner", 20.0, 0.35

    words = combined.split()
    technical = sum(1 for token in tokens if token in TECHNICAL_TERMS)
    density = technical * 100.0 / max(1, len(tokens))
    advanced_hits = sum(1 for token in tokens if token in ADVANCED_TERMS)
    beginner_hits = sum(
        1
        for term in BEGINNER_TERMS
        if re.search(rf"\b{re.escape(term)}\b", combined.lower())
    )

    sentences = _split_sentences(combined) or [combined]
    average_sentence = len(words) / max(1, len(sentences))
    average_word = sum(len(word) for word in words) / max(1, len(words))
    variety = len(set(tokens)) / max(1, len(tokens))

    score = 28.0
    score += min(34.0, density * 4.5)
    score += max(0.0, (average_word - 4.1) * 9.0)
    score += min(18.0, max(0.0, average_sentence - 12.0) * 0.9)
    score += min(12.0, advanced_hits * 2.0)
    score += max(-12.0, min(12.0, (variety - 0.45) * 25.0))
    score -= min(26.0, beginner_hits * 8.0)
    score = round(max(0.0, min(100.0, score)), 1)

    if score < 35:
        label = "beginner"
    elif score < 65:
        label = "intermediate"
    else:
        label = "advanced"

    distance = min(abs(score - 35), abs(score - 65))
    confidence = 0.55 + min(0.25, len(tokens) / 2500.0) + min(0.12, distance / 100.0)
    return label, score, round(min(0.95, confidence), 2)


def _analyze_text(
    text,
    *,
    title_hint=None,
    filename="",
    source="text",
    warnings=None,
    started_at=None,
):
    warnings = list(warnings or [])
    text = _cap_words(_clean_text(text))
    word_count = len(text.split())

    title, title_confidence = _guess_title(text, title_hint, filename)
    body = _strip_heading_lines(text, title)
    categories = list(Category.objects.all())
    category_name, category_id, category_confidence = _predict_category(
        body,
        title,
        [],
        categories,
    )
    tags = _extract_tags(body, title, category_name)
    category_name, category_id, category_confidence = _predict_category(
        body,
        title,
        tags,
        categories,
    )
    description = _summarize(body, title, tags)
    difficulty, difficulty_score, difficulty_confidence = _classify_difficulty(
        body, title
    )

    if not word_count:
        warnings.append(
            "No readable text was found, so suggestions rely on the available metadata."
        )

    processing_ms = None
    if started_at is not None:
        processing_ms = round((time.perf_counter() - started_at) * 1000, 2)

    return {
        "title": title,
        "description": description,
        "category": category_name,
        "category_id": category_id,
        "tags": tags,
        "difficulty": difficulty,
        "difficulty_score": difficulty_score,
        "confidence": {
            "title": round(title_confidence, 2),
            "category": category_confidence,
            "difficulty": difficulty_confidence,
        },
        "source": source,
        "word_count": word_count,
        "warnings": warnings,
        "processing_time_ms": processing_ms,
    }


def _extract_pdf(path):
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as error:
                raise ContentAnalysisError(
                    "This PDF is password protected.",
                    code="encrypted_file",
                ) from error

        pages = reader.pages[: _pdf_pages()]
        text_parts = []
        for page in pages:
            try:
                text_parts.append(page.extract_text() or "")
            except Exception:
                logger.warning("Unable to extract text from a PDF page", exc_info=True)

        metadata = reader.metadata or {}
        title_hint = None
        for key in ("/Title", "title"):
            value = metadata.get(key) if hasattr(metadata, "get") else None
            if value:
                title_hint = str(value).strip()
                break
        if title_hint and title_hint.lower() in {"untitled", "untitled document"}:
            title_hint = None
    except ContentAnalysisError:
        raise
    except Exception as error:
        raise ContentAnalysisError(
            "This PDF could not be read.",
            code="invalid_file",
        ) from error

    warnings = []
    text = _clean_text("\n".join(text_parts))
    if not text:
        warnings.append(
            "This PDF contains no selectable text. Scanned pages are not OCRed."
        )
    return ExtractedContent(text=text, title_hint=title_hint, source="pdf", warnings=warnings)


def _extract_docx(path):
    from docx import Document

    try:
        document = Document(str(path))
    except Exception as error:
        raise ContentAnalysisError(
            "This Word document could not be read.",
            code="invalid_file",
        ) from error

    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    title_hint = (document.core_properties.title or "").strip() or None
    text = _clean_text("\n".join(parts))
    warnings = [] if text else ["This document contains no readable text."]
    return ExtractedContent(
        text=text,
        title_hint=title_hint,
        source="docx",
        warnings=warnings,
    )


def _load_whisper_model():
    global _WHISPER_MODEL

    try:
        import whisper
    except ImportError as error:
        raise ContentAnalysisError(
            "Local speech recognition is not installed on this server.",
            code="transcription_unavailable",
        ) from error

    model_name = _setting(
        "CONTENT_ANALYZER_WHISPER_MODEL", DEFAULT_WHISPER_MODEL
    )
    with _WHISPER_LOCK:
        if _WHISPER_MODEL is None:
            logger.info("Loading local Whisper model %s", model_name)
            _WHISPER_MODEL = whisper.load_model(model_name, device="cpu")
    return _WHISPER_MODEL


def _read_wav_samples(path):
    # Whisper shells out to a system ffmpeg when given a file path; the
    # extracted WAV is already mono 16 kHz, so decode it ourselves instead.
    with wave.open(str(path), "rb") as wav_file:
        frames = wav_file.readframes(wav_file.getnframes())
    return np.frombuffer(frames, dtype=np.int16).astype(np.float32) / 32768.0


def transcribe_audio(audio_path):
    model = _load_whisper_model()
    language = _setting("CONTENT_ANALYZER_LANGUAGE", "en") or None
    result = model.transcribe(
        _read_wav_samples(audio_path),
        fp16=False,
        language=language,
        condition_on_previous_text=False,
    )
    return _clean_text(result.get("text") or "")


def _extract_media_audio(path, workdir):
    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
    destination = Path(workdir) / "analysis.wav"
    command = [
        ffmpeg,
        "-hide_banner",
        "-loglevel",
        "error",
        "-nostdin",
        "-y",
        "-t",
        str(_media_seconds()),
        "-i",
        str(path),
        "-vn",
        "-ac",
        "1",
        "-ar",
        "16000",
        "-f",
        "wav",
        str(destination),
    ]

    try:
        result = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=_setting("CONTENT_ANALYZER_FFMPEG_TIMEOUT_SECONDS", 300),
        )
    except subprocess.TimeoutExpired as error:
        raise ContentAnalysisError(
            "Audio extraction took too long. Try a shorter file.",
            code="audio_extraction_failed",
        ) from error
    except OSError as error:
        raise ContentAnalysisError(
            "The audio extractor could not be started.",
            code="audio_extraction_failed",
        ) from error

    if result.returncode != 0 or not destination.exists():
        logger.warning(
            "FFmpeg analysis extraction failed: %s", (result.stderr or "")[-2000:]
        )
        raise ContentAnalysisError(
            "The audio track could not be extracted from this file.",
            code="audio_extraction_failed",
        )
    return destination


def _extract_media(path, extension):
    source = "video" if extension in VIDEO_EXTENSIONS else "audio"
    with tempfile.TemporaryDirectory(prefix="edumind-audio-") as workdir:
        try:
            audio_path = _extract_media_audio(path, workdir)
        except ContentAnalysisError as error:
            return ExtractedContent(
                text="",
                source=source,
                warnings=[f"{error} Suggestions rely on the file name."],
            )

        try:
            text = transcribe_audio(audio_path)
        except ContentAnalysisError as error:
            return ExtractedContent(
                text="",
                source=source,
                warnings=[f"{error} Suggestions rely on the file name."],
            )
        except Exception:
            logger.exception("Local transcription failed")
            return ExtractedContent(
                text="",
                source=source,
                warnings=[
                    "Speech recognition failed for this file. "
                    "Suggestions rely on the file name."
                ],
            )

    warnings = []
    if not text:
        warnings.append(
            "No speech was detected in the analyzed portion of this file."
        )
    return ExtractedContent(text=text, source=source, warnings=warnings)


def _analyze_path(path, extension, *, filename, title_hint=None, source=None):
    if extension in PDF_EXTENSIONS:
        extracted = _extract_pdf(path)
    elif extension == "docx":
        extracted = _extract_docx(path)
    elif extension in AUDIO_EXTENSIONS or extension in VIDEO_EXTENSIONS:
        extracted = _extract_media(path, extension)
    else:
        raise ContentAnalysisError(
            "Unsupported file type. Use a PDF, DOCX, audio or video file.",
            code="unsupported_file",
        )

    if title_hint:
        extracted.title_hint = title_hint
    if source:
        extracted.source = source
    extracted.word_count = len(extracted.text.split())
    return extracted


def validate_analysis_upload(upload):
    filename = getattr(upload, "name", "") or ""
    extension = Path(filename).suffix.lower().lstrip(".")
    if extension not in ANALYZABLE_EXTENSIONS:
        raise ContentAnalysisError(
            "Unsupported file type. Use a PDF, DOCX, audio or video file.",
            code="unsupported_file",
        )

    size = getattr(upload, "size", 0) or 0
    if size <= 0:
        raise ContentAnalysisError("The file is empty.", code="empty_file")

    limit = (
        settings.MAX_VIDEO_UPLOAD_SIZE_BYTES
        if extension in VIDEO_EXTENSIONS
        else settings.MAX_UPLOAD_SIZE_BYTES
    )
    if size > limit:
        limit_mb = (
            settings.MAX_VIDEO_UPLOAD_SIZE_MB
            if extension in VIDEO_EXTENSIONS
            else settings.MAX_UPLOAD_SIZE_MB
        )
        raise ContentAnalysisError(
            f"The file exceeds the maximum allowed size of {limit_mb} MB.",
            code="file_too_large",
        )

    if hasattr(upload, "seek"):
        upload.seek(0)

    if extension == "docx":
        if not zipfile.is_zipfile(upload):
            raise ContentAnalysisError(
                "This Word document is not a valid DOCX file.",
                code="invalid_file",
            )
        try:
            upload.seek(0)
            with zipfile.ZipFile(upload) as archive:
                names = set(archive.namelist())
                uncompressed = sum(info.file_size for info in archive.infolist())
                if "word/document.xml" not in names:
                    raise ContentAnalysisError(
                        "This file is not a valid Word document.",
                        code="invalid_file",
                    )
                if uncompressed > DOCX_MAX_UNCOMPRESSED_BYTES:
                    raise ContentAnalysisError(
                        "This Word document is too large to analyze.",
                        code="file_too_large",
                    )
        except zipfile.BadZipFile as error:
            raise ContentAnalysisError(
                "This Word document is not a valid DOCX file.",
                code="invalid_file",
            ) from error
        finally:
            upload.seek(0)
        return

    try:
        validate_course_file(upload)
    except DjangoValidationError as error:
        raise ContentAnalysisError(
            " ".join(error.messages), code="invalid_file"
        ) from error
    finally:
        upload.seek(0)


def get_media_duration_seconds(path):
    """Return the real duration of an audio/video file, or None.

    imageio-ffmpeg only bundles the ffmpeg binary, so the duration is read
    from ffmpeg's own probe output instead of ffprobe.
    """
    command = [
        imageio_ffmpeg.get_ffmpeg_exe(),
        "-hide_banner",
        "-nostdin",
        "-i",
        str(path),
    ]
    try:
        result = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=_setting("CONTENT_ANALYZER_FFMPEG_TIMEOUT_SECONDS", 300),
        )
    except (subprocess.TimeoutExpired, OSError):
        logger.warning("Unable to probe media duration", exc_info=True)
        return None

    match = MEDIA_DURATION_PATTERN.search(result.stderr or "")
    if match is None:
        return None

    hours, minutes, seconds = match.groups()
    return int(hours) * 3600 + int(minutes) * 60 + float(seconds)


def estimate_lesson_duration(file_path, extension, word_count=0):
    """Estimate a lesson length in minutes.

    PDFs and documents use a reading speed of 200 words per minute; audio and
    video use the real media duration. Returns None when it cannot be known.
    """
    if extension in AUDIO_EXTENSIONS or extension in VIDEO_EXTENSIONS:
        seconds = get_media_duration_seconds(file_path)
        if not seconds:
            return None
        return max(1, round(seconds / 60))

    if word_count:
        return max(1, round(word_count / READING_WORDS_PER_MINUTE))
    return None


def _analyze_upload(upload):
    started_at = time.perf_counter()
    filename = getattr(upload, "name", "") or ""
    extension = Path(filename).suffix.lower().lstrip(".")
    validate_analysis_upload(upload)

    with tempfile.TemporaryDirectory(prefix="edumind-analysis-") as workdir:
        path = Path(workdir) / f"source.{extension}"
        try:
            upload.seek(0)
            with path.open("wb") as target:
                shutil.copyfileobj(upload, target)
        except OSError as error:
            raise ContentAnalysisError(
                "The uploaded file could not be read.",
                code="invalid_file",
            ) from error

        extracted = _analyze_path(path, extension, filename=filename)
        duration = estimate_lesson_duration(
            path, extension, extracted.word_count
        )

    result = _analyze_text(
        extracted.text,
        title_hint=extracted.title_hint,
        filename=filename,
        source=extracted.source,
        warnings=extracted.warnings,
        started_at=started_at,
    )
    return result, duration


def analyze_upload(upload):
    result, _duration = _analyze_upload(upload)
    return result


def _lesson_sort_key(lesson, index):
    stem = Path(lesson["file_name"] or "").stem
    prefix = FILENAME_NUMBER_PREFIX.match(stem)
    return (
        DIFFICULTY_ORDER.get(lesson["difficulty"], len(DIFFICULTY_ORDER)),
        0 if prefix else 1,
        int(prefix.group(1)) if prefix else 0,
        stem.lower(),
        index,
    )


def order_lessons(analyzed_lessons):
    """Order lessons from beginner to advanced, then by file name.

    Files with a number prefix ("01_intro") keep their intended sequence;
    everything else falls back to the alphabetical file name and finally to
    the original upload order.
    """
    ordered = sorted(
        enumerate(analyzed_lessons),
        key=lambda item: _lesson_sort_key(item[1], item[0]),
    )
    lessons = []
    for position, (_index, lesson) in enumerate(ordered, start=1):
        lessons.append({**lesson, "order": position})
    return lessons


def analyze_bulk_lessons(files):
    """Analyze several uploads at once, skipping the ones that fail."""
    started_at = time.perf_counter()
    analyzed = []
    failed = []

    for upload in files:
        filename = getattr(upload, "name", "") or ""
        try:
            result, duration = _analyze_upload(upload)
        except ContentAnalysisError as error:
            failed.append({"file_name": filename, "error": str(error)})
            continue
        except Exception:
            logger.exception("Bulk lesson analysis failed for %s", filename)
            failed.append(
                {
                    "file_name": filename,
                    "error": "This file could not be analyzed.",
                }
            )
            continue

        warnings = list(result["warnings"])
        if duration is None:
            warnings.append("The lesson duration could not be estimated.")

        confidence = (
            result["confidence"]["title"] + result["confidence"]["difficulty"]
        ) / 2
        analyzed.append(
            {
                "title": result["title"],
                "description": result["description"],
                "duration_minutes": duration,
                "file_name": filename,
                "file_type": result["source"],
                "difficulty": result["difficulty"],
                "confidence": round(confidence, 2),
                "warnings": warnings,
            }
        )

    lessons = order_lessons(analyzed)
    return {
        "lessons": lessons,
        "processing_time_ms": round(
            (time.perf_counter() - started_at) * 1000, 2
        ),
        "total_files": len(files),
        "successful": len(lessons),
        "failed": failed,
    }


def analyze_course(course):
    started_at = time.perf_counter()
    transcript = getattr(course, "transcript", None)
    if (
        transcript is not None
        and transcript.status == CourseTranscript.Status.COMPLETED
        and transcript.text.strip()
    ):
        return _analyze_text(
            transcript.text,
            title_hint=course.title,
            filename=course.file.name if course.file else "",
            source="transcript",
            warnings=[],
            started_at=started_at,
        )

    # Courses are lesson-based now: fall back to the first lesson file when
    # the legacy course-level material is missing.
    material = course.file
    if not material:
        first_lesson = (
            course.lessons.exclude(file="").order_by("order").first()
        )
        if first_lesson is not None:
            material = first_lesson.file

    if not material:
        text = " ".join(
            part
            for part in (course.title, course.description, course.tags)
            if part
        )
        return _analyze_text(
            text,
            title_hint=course.title,
            source="metadata",
            warnings=[
                "This course has no uploaded material. Suggestions use its "
                "existing details."
            ],
            started_at=started_at,
        )

    extension = Path(material.name).suffix.lower().lstrip(".")
    if extension not in ANALYZABLE_EXTENSIONS:
        raise ContentAnalysisError(
            "This course file cannot be analyzed.",
            code="unsupported_file",
        )

    with tempfile.TemporaryDirectory(prefix="edumind-analysis-") as workdir:
        path = Path(workdir) / f"source.{extension}"
        try:
            with material.storage.open(material.name, "rb") as source:
                with path.open("wb") as target:
                    shutil.copyfileobj(source, target)
        except (FileNotFoundError, OSError) as error:
            raise ContentAnalysisError(
                "The course file could not be found.",
                code="missing_file",
            ) from error

        extracted = _analyze_path(
            path,
            extension,
            filename=material.name,
            title_hint=course.title,
        )

    return _analyze_text(
        extracted.text,
        title_hint=extracted.title_hint,
        filename=material.name,
        source=extracted.source,
        warnings=extracted.warnings,
        started_at=started_at,
    )
